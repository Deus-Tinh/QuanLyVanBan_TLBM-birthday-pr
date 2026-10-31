"""Xuất Excel danh bạ Phòng ban theo mẫu danh bạ: Khối/Khu vực -> Đơn vị -> Danh sách CBNV.

Quy tắc (chốt 15-18/10/2026, điều chỉnh 24/10/2026):
- Khối theo thu_tu_hien_thi, đơn vị theo stt_uu_tien; STT người đếm lại từ 1 ở mỗi đơn vị.
- Dưới mỗi đơn vị liệt kê mọi CBNV đang làm việc của đơn vị, chức vụ lớn trước (display_order nhỏ trước),
  bằng nhau thì id nhỏ trước. Đơn vị không có ai chỉ có dòng đơn vị.
- Tài khoản quản trị hệ thống (is_superuser) không phải CBNV của đơn vị nên không được liệt kê.
- Người kiêm nhiệm (accounts_employee_assignment) được liệt kê thêm ở mỗi đơn vị kiêm nhiệm với chức vụ
  kiêm nhiệm; đơn vị và chức vụ chính vẫn giữ nguyên. Cùng một người, đơn vị, chức vụ chỉ xuất một dòng.
- CBNV Nội bộ (is_public = False) chỉ xuất khi người xuất có view_full_cbnv hoặc là superuser.
- Ô văn bản bắt đầu bằng "=" luôn ghi dạng chuỗi để Excel không hiểu là công thức.
"""
import io

import openpyxl
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from accounts.models import Employee, EmployeeAssignment

COLUMN_WIDTHS = [7, 30, 26, 44, 16, 34]
COLUMN_COUNT = len(COLUMN_WIDTHS)
HEADING_FONT = "Times New Roman"
BODY_FONT = "Arial"
# Màu theo mẫu danh bạ: chữ khối xanh, nền dòng tiêu đề cột xanh nhạt.
BLOCK_COLOR = "3300FF"
HEADER_FILL = "99E6FF"
BLOCK_FILL = "EEF1FF"
UNIT_FILL = "F3F4F6"
BORDER_COLOR = "BFC5CD"
SUBTITLE_COLOR = "4B5563"
HEADER_ROW = 5


def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _set_text(cell, value):
    """Ghi giá trị; chuỗi bắt đầu bằng dấu bằng luôn là chuỗi, không phải công thức."""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def can_export_internal_people(user) -> bool:
    return bool(user.is_superuser or user.has_perm("view_full_cbnv"))


class Member:
    """Một dòng người trong đơn vị: đơn vị chính (Employee) hoặc đơn vị kiêm nhiệm (EmployeeAssignment)."""

    def __init__(self, employee, position, title):
        self.employee = employee
        self.position = position
        self.title = title

    @property
    def sort_key(self):
        order = self.position.display_order if self.position else float("inf")
        return (order, self.employee.id)


def find_unit_members(unit_ids, include_internal_people):
    """Trả về {don_vi_id: [Member, ...]} CBNV từng đơn vị theo thứ tự chức vụ (đơn vị chính và kiêm nhiệm).

    Hai truy vấn cho mọi đơn vị, không phụ thuộc số đơn vị.
    """
    employee_filter = {"trang_thai": Employee.Status.LAM_VIEC, "is_superuser": False}
    if not include_internal_people:
        employee_filter["is_public"] = True
    primary = Employee.objects.filter(don_vi_id__in=unit_ids, **employee_filter).select_related("chuc_vu")
    concurrent = EmployeeAssignment.objects.filter(
        don_vi_id__in=unit_ids, **{f"employee__{key}": value for key, value in employee_filter.items()}
    ).select_related("employee", "chuc_vu")

    members = {}
    seen = set()
    for employee in primary:
        members.setdefault(employee.don_vi_id, []).append(Member(employee, employee.chuc_vu, employee.chuc_danh))
        seen.add((employee.id, employee.don_vi_id, employee.chuc_vu_id))
    for assignment in concurrent:
        key = (assignment.employee_id, assignment.don_vi_id, assignment.chuc_vu_id)
        if key in seen:
            continue
        seen.add(key)
        members.setdefault(assignment.don_vi_id, []).append(
            Member(assignment.employee, assignment.chuc_vu, assignment.chuc_danh)
        )
    for rows in members.values():
        rows.sort(key=lambda member: member.sort_key)
    return members


def unit_label(unit, i18n):
    """Dòng đơn vị như mẫu: TÊN ĐƠN VỊ, sau đó địa chỉ, điện thoại, email nếu có."""
    parts = [unit.ten.upper()]
    if unit.dia_chi:
        parts.append(f"{i18n['address_label']}: {unit.dia_chi}")
    if unit.dien_thoai:
        parts.append(f"{i18n['phone_label']}: {unit.dien_thoai}")
    if unit.email:
        parts.append(f"{i18n['email_label']}: {unit.email}")
    return "  |  ".join(parts)


def build_directory_workbook(units, user, i18n):
    """`units`: queryset DonVi đã lọc theo quyền và bộ lọc. Trả về (nội dung .xlsx, thời điểm xuất)."""
    units = list(
        units.select_related("khoi").order_by(
            "khoi__thu_tu_hien_thi", "khoi__ten", "khoi_id", "stt_uu_tien", "ten", "id"
        )
    )
    members = find_unit_members([unit.id for unit in units], can_export_internal_people(user))
    generated_at = timezone.localtime()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = i18n["sheet_title"]
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def merged_row(row, text, font, alignment, fill=None, height=None, boxed=True):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COLUMN_COUNT)
        cell = ws.cell(row=row, column=1)
        _set_text(cell, text)
        cell.font = font
        cell.alignment = alignment
        for col in range(1, COLUMN_COUNT + 1):
            target = ws.cell(row=row, column=col)
            if boxed:
                target.border = border
            if fill:
                target.fill = _fill(fill)
        if height:
            ws.row_dimensions[row].height = height

    merged_row(1, i18n["title"], Font(name=HEADING_FONT, size=14, bold=True, color=BLOCK_COLOR), center,
               height=26, boxed=False)
    merged_row(2, i18n["subtitle"].upper(), Font(name=HEADING_FONT, size=12, bold=True, color=BLOCK_COLOR),
               center, height=22, boxed=False)
    merged_row(3, f"{i18n['generated_label']}: {generated_at.strftime('%d/%m/%Y %H:%M')}",
               Font(name=BODY_FONT, size=10, italic=True, color=SUBTITLE_COLOR), center, boxed=False)

    for col, label in enumerate(i18n["headers"], start=1):
        cell = ws.cell(row=HEADER_ROW, column=col)
        cell.value = label
        cell.font = Font(name=BODY_FONT, size=11, bold=True)
        cell.fill = _fill(HEADER_FILL)
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[HEADER_ROW].height = 22

    row = HEADER_ROW + 1
    current_block_id = object()
    for unit in units:
        if unit.khoi_id != current_block_id:
            current_block_id = unit.khoi_id
            merged_row(row, unit.khoi.ten.upper(), Font(name=HEADING_FONT, size=12, bold=True, color=BLOCK_COLOR),
                       center, BLOCK_FILL, height=22)
            row += 1
        merged_row(row, unit_label(unit, i18n), Font(name=HEADING_FONT, size=11, bold=True), left, UNIT_FILL,
                   height=20)
        row += 1
        for number, member in enumerate(members.get(unit.id, []), start=1):
            values = [
                number,
                member.employee.ho_ten,
                member.position.name if member.position else "",
                member.title or "",
                member.employee.sdt or "",
                member.employee.email or "",
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col)
                _set_text(cell, value)
                cell.font = Font(name=BODY_FONT, size=11)
                cell.border = border
                cell.alignment = center if col in (1, 5) else left
            row += 1

    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), generated_at
